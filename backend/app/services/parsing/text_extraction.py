"""Safe text extraction from uploaded PDF and DOCX files.

Every failure is converted to a DocumentError with a user-facing message; malformed files never
propagate library exceptions to the API layer.
"""

import io
import logging
import zipfile
from dataclasses import dataclass
from typing import Literal

import docx
import pymupdf

from app.services.parsing.text_utils import normalize_text

logger = logging.getLogger(__name__)

DocumentKind = Literal["pdf", "docx"]
MIN_TEXT_CHARS = 150
MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_DOCX_ENTRIES = 2000

CONTENT_TYPES: dict[DocumentKind, str] = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


class DocumentError(Exception):
    def __init__(self, message: str, code: str = "invalid_document") -> None:
        super().__init__(message)
        self.message = message
        self.code = code


@dataclass
class ExtractedDocument:
    kind: DocumentKind
    text: str
    page_count: int | None


def detect_kind(filename: str, content: bytes) -> DocumentKind:
    """Identify the file by its magic bytes and require the extension to agree."""
    name = filename.lower()
    if content.startswith(b"%PDF-"):
        if not name.endswith(".pdf"):
            raise DocumentError("The file extension does not match its contents.", "file_type_mismatch")
        return "pdf"
    if content.startswith(b"PK\x03\x04"):
        if not name.endswith(".docx"):
            raise DocumentError("Only .docx Word documents are supported.", "unsupported_file_type")
        return "docx"
    if name.endswith(".doc"):
        raise DocumentError(
            "Legacy .doc files are not supported. Please save the document as .docx or PDF.", "unsupported_file_type"
        )
    raise DocumentError("Unsupported file type. Please upload a PDF or DOCX resume.", "unsupported_file_type")


def _extract_pdf(content: bytes, max_pages: int) -> ExtractedDocument:
    try:
        doc = pymupdf.open(stream=content, filetype="pdf")
    except Exception as exc:  # PyMuPDF raises several unrelated exception types for corrupt input
        logger.info("PDF open failed: %s", type(exc).__name__)
        raise DocumentError("The PDF could not be read. It may be corrupted.") from exc
    with doc:
        if doc.needs_pass or doc.is_encrypted:
            raise DocumentError("Password-protected PDFs are not supported.", "encrypted_document")
        if doc.page_count == 0:
            raise DocumentError("The PDF has no pages.", "empty_document")
        if doc.page_count > max_pages:
            raise DocumentError(f"Resumes longer than {max_pages} pages are not supported.", "too_many_pages")
        try:
            pages = [page.get_text("text", sort=True) for page in doc]
        except Exception as exc:
            logger.info("PDF text extraction failed: %s", type(exc).__name__)
            raise DocumentError("Text could not be extracted from this PDF.") from exc
        page_count = doc.page_count
    text = normalize_text("\n".join(pages))
    if len(text) < MIN_TEXT_CHARS:
        raise DocumentError(
            "Little or no selectable text was found. This looks like a scanned or image-only PDF, and OCR is "
            "not supported. Please upload a text-based PDF or a DOCX file.",
            "no_text_found",
        )
    return ExtractedDocument("pdf", text, page_count)


def _check_zip(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            infos = zf.infolist()
            if len(infos) > MAX_DOCX_ENTRIES or sum(i.file_size for i in infos) > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise DocumentError("The document is too large when decompressed.", "invalid_document")
            if "word/document.xml" not in zf.namelist():
                raise DocumentError("The file is not a valid Word document.")
    except zipfile.BadZipFile as exc:
        raise DocumentError("The DOCX file is corrupted.") from exc


def _extract_docx(content: bytes) -> ExtractedDocument:
    _check_zip(content)
    try:
        document = docx.Document(io.BytesIO(content))
    except Exception as exc:
        logger.info("DOCX open failed: %s", type(exc).__name__)
        raise DocumentError("The DOCX file could not be read.") from exc
    lines: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name if para.style is not None else "") or ""
        # Word list paragraphs carry no bullet glyph in their text, so restore one for the parser.
        if "List" in style and not text.startswith(("•", "-")):
            text = f"• {text}"
        lines.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            lines.extend(dict.fromkeys(cells))  # merged cells repeat their text
    text = normalize_text("\n".join(lines))
    if len(text) < MIN_TEXT_CHARS:
        raise DocumentError("The document contains too little text to analyze.", "no_text_found")
    return ExtractedDocument("docx", text, None)


def extract_text(filename: str, content: bytes, *, max_pages: int) -> ExtractedDocument:
    if not content:
        raise DocumentError("The uploaded file is empty.", "empty_document")
    kind = detect_kind(filename, content)
    return _extract_pdf(content, max_pages) if kind == "pdf" else _extract_docx(content)
