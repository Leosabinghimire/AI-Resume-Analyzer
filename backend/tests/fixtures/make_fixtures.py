"""Generate the sample resume fixtures used by tests and the demo seed.

The person and companies are fictional. Run:  python tests/fixtures/make_fixtures.py
"""

from pathlib import Path

import docx
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen import canvas

HERE = Path(__file__).parent

SAMPLE_LINES = [
    "Alex Rivera",
    "alex.rivera@example.com | +1 555 010 7788 | Austin, TX | github.com/alexrivera-demo",
    "",
    "SUMMARY",
    "Full-stack software engineer with 5 years of experience building web applications with React,",
    "TypeScript and Python. Focused on reliable APIs, clean UI and automated testing.",
    "",
    "SKILLS",
    "Languages: Python, TypeScript, JavaScript, SQL",
    "Frameworks: React, Next.js, FastAPI, Django",
    "Tools: PostgreSQL, Docker, Git, GitHub Actions, AWS",
    "",
    "EXPERIENCE",
    "Senior Software Engineer | Northwind Analytics",
    "Jan 2022 - Present",
    "• Built REST APIs with FastAPI and PostgreSQL serving 40 internal dashboards",
    "• Led migration of the frontend from JavaScript to TypeScript and React with a team of 4 engineers",
    "• Responsible for code reviews and mentoring two junior developers",
    "• Set up CI/CD pipelines with GitHub Actions, reducing release time by 30%",
    "Software Engineer | Blue Harbor Labs",
    "Jun 2019 - Dec 2021",
    "• Worked on Django services for customer onboarding",
    "• Developed reusable React components used across 3 products",
    "• Helped with database performance tuning in PostgreSQL",
    "",
    "PROJECTS",
    "Recipe Planner - personal project",
    "• Next.js and TypeScript app for meal planning with a FastAPI backend",
    "",
    "EDUCATION",
    "Bachelor of Science in Computer Science, University of Texas",
    "2015 - 2019",
    "",
    "CERTIFICATIONS",
    "AWS Certified Developer - Amazon Web Services, 2023",
]

SAMPLE_JOB = """Senior Full-Stack Engineer

About us
We build analytics software for logistics teams.

Responsibilities
- Design, build and maintain REST APIs using Python and FastAPI
- Build responsive user interfaces with React and TypeScript
- Write unit and integration tests and participate in code reviews
- Improve CI/CD pipelines and deployment automation on AWS

Requirements
- 4+ years of professional software development experience
- Strong experience with Python, React and TypeScript
- Experience with PostgreSQL and Docker
- Bachelor's degree in Computer Science or equivalent experience

Nice to have
- Experience with Kubernetes and GraphQL
- Familiarity with Terraform
"""


def make_pdf(path: Path) -> None:
    c = canvas.Canvas(str(path), pagesize=LETTER)
    y = 750
    for line in SAMPLE_LINES:
        if line:
            c.setFont("Helvetica-Bold" if line.isupper() or line == SAMPLE_LINES[0] else "Helvetica", 10)
            c.drawString(50, y, line)
        y -= 14
    c.save()


def make_docx(path: Path) -> None:
    document = docx.Document()
    for line in SAMPLE_LINES:
        if not line:
            continue
        if line.startswith("• "):
            document.add_paragraph(line[2:], style="List Bullet")
        else:
            document.add_paragraph(line)
    document.save(str(path))


if __name__ == "__main__":
    make_pdf(HERE / "sample_resume.pdf")
    make_docx(HERE / "sample_resume.docx")
    (HERE / "sample_job.txt").write_text(SAMPLE_JOB, encoding="utf-8")
    print("Fixtures written to", HERE)
